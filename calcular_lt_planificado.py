#!/usr/bin/env python3
"""
Cálculo del LT planificado por puesto de trabajo.

Lee la hoja 'Detalle' de la exportación del ERP (una fila por puesto y día con el nº de
órdenes en cola y los días laborables que lleva esperando la más antigua) y genera
LT_planificado_AAAA-MM.xlsx con el LT a congelar por puesto, los avisos y la evolución mensual.

Uso:
    python calcular_lt_planificado.py [excel_entrada] [--vigente LT_planificado_AAAA-MM.xlsx]

Con --vigente se compara con el LT congelado del mes anterior (columna 'LT propuesto' de ese
fichero, que se puede corregir a mano) y solo se propone cambiar cuando la diferencia supera
la banda muerta.
"""
import argparse
import datetime as dt
import math
import statistics
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Parámetros
# ---------------------------------------------------------------------------
VENTANA_MESES = 6            # meses naturales usados para el LT (incluido el mes de corte)
VENTANA_RECIENTE_MESES = 3   # meses para detectar cambios recientes
MIN_SALIDAS = 8              # por debajo: aviso POCOS DATOS

# NO FIABLE automático: casi no salen órdenes y la más antigua lleva mucho tiempo
NO_FIABLE_MAX_SALIDAS = 5
NO_FIABLE_ANTIGUEDAD = 30
# Puestos que nunca se calculan (administrativos u órdenes que no se cierran en el ERP)
NO_FIABLE_MANUAL = {
    "DOCUMENTACION",
    "ASRM-CARMONA",
    "OPERACIONES-AUXILIARES-CURADO",
    "UTILLAJE",
    "VERIF-UTILLAJE",
}

BLOQUEO_MIN_DIAS = 10        # BLOQUEO ACTUAL si la más antigua lleva > max(10, 3 x LT)
BLOQUEO_FACTOR = 3
CAMBIO_MIN_DIAS = 3          # CAMBIO RECIENTE si |P50 3 m - P50 6 m| > max(3, 50 %)
CAMBIO_PCT = 0.50
CAMBIO_MIN_SALIDAS = 4
WIP_FACTOR = 1.5             # WIP CRECIENTE si WIP últimos 30 d > 1,5 x WIP medio de la ventana
WIP_MIN = 10
WIP_DIAS_RECIENTES = 30

BANDA_PCT = 0.20             # banda muerta: solo se cambia si la diferencia es >= 20 %...
BANDA_MIN_DIAS = 1           # ...y >= 1 día laborable
TOPE_CAMBIO_PCT = 0.30       # cambio máximo por revisión

# Ruta tipo para el ejemplo del colchón en la hoja 'Criterio y revisión'
RUTA_EJEMPLO = [
    "CORTE-TELAS", "LAY-UP", "AUTOCLAVES", "BERMAQ-5X-12M", "VERIF-RECANTEO",
    "PPFF-AUTO", "PINTURA-AUTO", "VERIF-PINTURA", "IDENT", "LOGISTICA-ALMACEN/ENVIO",
]

MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]


# ---------------------------------------------------------------------------
# Lectura
# ---------------------------------------------------------------------------
def a_fecha(valor):
    """Convierte la celda FECHA (fecha, número de serie de Excel o texto) a date."""
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    if isinstance(valor, (int, float)):
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(valor))
    texto = str(valor).strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    raise ValueError(f"Fecha no reconocida: {valor!r}")


def leer_detalle(libro):
    """Devuelve {puesto: [(fecha, wip, antigüedad), ...]} ordenado por fecha y {puesto: área}."""
    hoja = libro["Detalle"]
    filas = hoja.iter_rows(values_only=True)
    cabecera = [str(c).strip().upper() if c is not None else "" for c in next(filas)]
    col = {nombre: cabecera.index(nombre) for nombre in ("PROVEEDOR", "AREA", "LANZAMIENTOS", "LT", "FECHA")}

    series, areas = defaultdict(list), {}
    for fila in filas:
        puesto = fila[col["PROVEEDOR"]]
        if puesto is None or fila[col["FECHA"]] is None:
            continue
        puesto = str(puesto).strip()
        series[puesto].append((
            a_fecha(fila[col["FECHA"]]),
            float(fila[col["LANZAMIENTOS"]] or 0),
            float(fila[col["LT"]] or 0),
        ))
        areas[puesto] = str(fila[col["AREA"]] or "").strip()
    for serie in series.values():
        serie.sort()
    return series, areas


def leer_resumen_q3(libro):
    """Devuelve {(puesto, año, mes): LT_Q3} de la hoja 'Resumen' (por posición de columna)."""
    if "Resumen" not in libro.sheetnames:
        return {}
    q3 = {}
    for fila in libro["Resumen"].iter_rows(min_row=2, values_only=True):
        if fila[2] is None or fila[5] is None:
            continue
        q3[(str(fila[2]).strip(), int(fila[0]), int(fila[1]))] = float(fila[5])
    return q3


def leer_vigente(ruta):
    """Lee el LT congelado del fichero del mes anterior (columna 'LT propuesto')."""
    hoja = load_workbook(ruta, read_only=True, data_only=True)["Listado"]
    filas = hoja.iter_rows(values_only=True)
    cabecera = [str(c or "") for c in next(filas)]
    i_puesto = cabecera.index("Puesto")
    i_lt = next(i for i, c in enumerate(cabecera) if c.startswith("LT propuesto"))
    return {
        str(fila[i_puesto]).strip(): int(round(float(fila[i_lt])))
        for fila in filas
        if fila[i_puesto] is not None and fila[i_lt] not in (None, "")
    }


# ---------------------------------------------------------------------------
# Cálculo
# ---------------------------------------------------------------------------
def percentil(valores, p):
    """Percentil con interpolación lineal (igual que PERCENTIL.INC de Excel)."""
    if not valores:
        return None
    v = sorted(valores)
    k = (len(v) - 1) * p
    i = int(k)
    j = min(i + 1, len(v) - 1)
    return v[i] + (v[j] - v[i]) * (k - i)


def media(valores):
    return statistics.mean(valores) if valores else None


def primer_dia_mes(fecha, meses_atras):
    """Primer día del mes situado 'meses_atras' meses antes del mes de 'fecha'."""
    total = fecha.year * 12 + fecha.month - 1 - meses_atras
    return dt.date(total // 12, total % 12 + 1, 1)


def salidas(serie):
    """Cada caída de la antigüedad diaria indica que ha salido la orden más antigua.
    El valor anterior a la caída es su LT real en el puesto (resolución ±1 día).
    Devuelve [(fecha de la caída, LT de la orden que sale), ...]."""
    return [(b[0], a[2]) for a, b in zip(serie, serie[1:]) if b[2] < a[2]]


def calcular_puesto(puesto, serie, corte, ini6, ini3, q3_resumen):
    sal = salidas(serie)
    s6 = [v for f, v in sal if f >= ini6]
    s3 = [v for f, v in sal if f >= ini3]
    p50_6, p75_6, p90_6 = (percentil(s6, p) for p in (0.50, 0.75, 0.90))
    p50_3 = percentil(s3, 0.50)

    antig_p50 = percentil([lt for f, _, lt in serie if f >= ini6], 0.50) or 0
    wip_6m = media([w for f, w, _ in serie if f >= ini6])
    wip_3m = media([w for f, w, _ in serie if f >= ini3])
    wip_30 = media([w for f, w, _ in serie if f > corte - dt.timedelta(days=WIP_DIAS_RECIENTES)])
    actual = serie[-1][2] if serie[-1][0] == corte else 0

    q3 = [v for (p, a, m), v in q3_resumen.items() if p == puesto and dt.date(a, m, 1) >= ini3]

    no_fiable = puesto in NO_FIABLE_MANUAL or (
        len(s6) < NO_FIABLE_MAX_SALIDAS and antig_p50 > NO_FIABLE_ANTIGUEDAD
    )
    lt = None if no_fiable or not s6 else max(1, math.ceil(p50_6))
    proteccion = None if lt is None else max(0, math.ceil(p75_6) - lt)

    avisos, acciones = [], []
    if no_fiable:
        avisos.append("NO FIABLE")
        acciones.append("No usar el dato: revisar en el ERP las órdenes abiertas que no se cierran "
                        "y fijar un LT manual con el responsable del puesto.")
    elif not s6:
        avisos.append("POCOS DATOS")
        acciones.append("Sin salidas en la ventana: fijar un LT manual con el jefe de área.")
    elif len(s6) < MIN_SALIDAS:
        avisos.append("POCOS DATOS")
        acciones.append("Valor orientativo (pocas salidas): validarlo con el jefe de área.")
    if lt is not None and actual > max(BLOQUEO_MIN_DIAS, BLOQUEO_FACTOR * lt):
        avisos.append(f"BLOQUEO ACTUAL ({actual:.0f} d)")
        acciones.append(f"La orden más antigua lleva {actual:.0f} días: revisar si está retenida "
                        "(material, NC, plano). No subir el LT por este motivo.")
    if (not no_fiable and len(s3) >= CAMBIO_MIN_SALIDAS and p50_6 is not None
            and abs(p50_3 - p50_6) > max(CAMBIO_MIN_DIAS, CAMBIO_PCT * p50_6)):
        avisos.append(f"CAMBIO RECIENTE (3 m: {p50_3:.0f} d)")
        acciones.append("El LT de los últimos 3 meses se aleja del de 6 meses: vigilar en la próxima revisión.")
    if (not no_fiable and wip_30 is not None and wip_6m
            and wip_30 >= WIP_MIN and wip_30 > WIP_FACTOR * wip_6m):
        avisos.append(f"WIP CRECIENTE (x{wip_30 / wip_6m:.1f})")
        acciones.append("La cola está creciendo: si no se actúa sobre la capacidad, el LT real subirá.")

    return {
        "puesto": puesto, "lt": lt, "proteccion": proteccion,
        "n6": len(s6), "p50_6": p50_6, "p75_6": p75_6, "p90_6": p90_6, "p50_3": p50_3,
        "wip_3m": wip_3m, "wip_30": wip_30, "actual": actual,
        "q3_resumen": media(q3),
        "aviso": " | ".join(avisos), "accion": " ".join(acciones),
        "salidas": sal, "serie": serie,
    }


def aplicar_banda(calculado, vigente):
    """Devuelve (LT propuesto, ¿cambiar?) aplicando la banda muerta y el tope por revisión."""
    if vigente is None:
        return calculado, "NUEVO" if calculado is not None else "FIJAR A MANO"
    if calculado is None:
        return vigente, "NO (sin dato fiable)"
    diferencia = calculado - vigente
    if abs(diferencia) < max(BANDA_MIN_DIAS, BANDA_PCT * vigente):
        return vigente, "NO"
    paso_max = max(1, round(TOPE_CAMBIO_PCT * vigente))
    propuesto = vigente + max(-paso_max, min(paso_max, diferencia))
    return propuesto, "SÍ" if propuesto == calculado else "SÍ (limitado al tope)"


# ---------------------------------------------------------------------------
# Escritura del Excel
# ---------------------------------------------------------------------------
NEGRITA = Font(bold=True)
BLANCO_NEGRITA = Font(bold=True, color="FFFFFF")
RELLENO_CABECERA = PatternFill("solid", fgColor="305496")
RELLENO_PROPUESTO = PatternFill("solid", fgColor="E2EFDA")
RELLENO_ROJO = PatternFill("solid", fgColor="FFC7CE")
RELLENO_NARANJA = PatternFill("solid", fgColor="FCE4D6")
RELLENO_AMARILLO = PatternFill("solid", fgColor="FFF2CC")


def cabecera(hoja, titulos, anchos):
    hoja.append(titulos)
    for i, ancho in enumerate(anchos, start=1):
        celda = hoja.cell(row=1, column=i)
        celda.font = BLANCO_NEGRITA
        celda.fill = RELLENO_CABECERA
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.row_dimensions[1].height = 45
    hoja.freeze_panes = "C2"


def hoja_listado(libro, resultados, areas, vigentes, hay_vigente):
    hoja = libro.active
    hoja.title = "Listado"
    titulos = [
        "Área", "Puesto", "LT calculado (P50, días lab.)", "Protección (P75−P50, días)",
        "LT vigente", "LT propuesto (a congelar)", "¿Cambiar?",
        "Nº salidas (6 m)", "P50 salidas 6 m", "P75 salidas 6 m", "P90 salidas 6 m", "P50 salidas 3 m",
        "WIP medio 3 m", "WIP últimos 30 d", "Antigüedad actual más antigua",
        "Q3 hoja Resumen (3 m)", "Aviso", "Acción sugerida",
    ]
    anchos = [14, 36, 12, 12, 10, 12, 14, 10, 10, 10, 10, 10, 10, 10, 12, 12, 34, 70]
    cabecera(hoja, titulos, anchos)

    for r in resultados:
        vigente = vigentes.get(r["puesto"])
        if hay_vigente:
            propuesto, cambiar = aplicar_banda(r["lt"], vigente)
        else:
            propuesto, cambiar = r["lt"], ""
        hoja.append([
            areas[r["puesto"]], r["puesto"], r["lt"], r["proteccion"],
            vigente, propuesto, cambiar,
            r["n6"], r["p50_6"], r["p75_6"], r["p90_6"], r["p50_3"],
            r["wip_3m"], r["wip_30"], r["actual"],
            r["q3_resumen"], r["aviso"], r["accion"],
        ])
        fila = hoja.max_row
        hoja.cell(row=fila, column=6).fill = RELLENO_PROPUESTO
        hoja.cell(row=fila, column=6).font = NEGRITA
        for c in (9, 10, 11, 12, 16):
            hoja.cell(row=fila, column=c).number_format = "0.0"
        for c in (13, 14, 15):
            hoja.cell(row=fila, column=c).number_format = "0"
        hoja.cell(row=fila, column=18).alignment = Alignment(wrap_text=True, vertical="top")
        if r["aviso"]:
            relleno = (RELLENO_ROJO if "NO FIABLE" in r["aviso"]
                       else RELLENO_NARANJA if "BLOQUEO" in r["aviso"] else RELLENO_AMARILLO)
            hoja.cell(row=fila, column=17).fill = relleno
            hoja.cell(row=fila, column=18).fill = relleno
    hoja.auto_filter.ref = hoja.dimensions


def hoja_criterio(libro, resultados, corte, ini6, ini3, ruta_vigente):
    hoja = libro.create_sheet("Criterio y revisión")
    hoja.column_dimensions["A"].width = 110
    hoja.column_dimensions["B"].width = 10
    hoja.column_dimensions["C"].width = 12
    por_puesto = {r["puesto"]: r for r in resultados}

    def linea(texto="", estilo=None, *extra):
        hoja.append([texto, *extra])
        celda = hoja.cell(row=hoja.max_row, column=1)
        celda.alignment = Alignment(wrap_text=True, vertical="top")
        if estilo == "titulo":
            celda.font = Font(bold=True, size=14)
        elif estilo == "seccion":
            celda.font = Font(bold=True, size=12, color="305496")
        elif estilo == "negrita":
            for c in range(1, 2 + len(extra)):
                hoja.cell(row=hoja.max_row, column=c).font = NEGRITA

    linea("Criterio de cálculo del LT planificado y reglas de revisión", "titulo")
    linea(f"Fecha de corte de los datos: {corte:%d/%m/%Y}. Ventana del LT: {ini6:%d/%m/%Y} – {corte:%d/%m/%Y} "
          f"({VENTANA_MESES} meses). Ventana reciente: desde {ini3:%d/%m/%Y}. "
          f"Fichero vigente usado para la banda muerta: {ruta_vigente or 'ninguno (primer cálculo)'}.")
    linea()

    linea("1. Qué contiene la exportación de origen", "seccion")
    linea("La hoja 'Detalle' es una foto diaria de la cola de cada puesto, no el lead time de órdenes terminadas: "
          "'Lanzamientos' es el nº de órdenes en cola ese día (WIP) y 'LT' son los días laborables que lleva "
          "esperando la orden más antigua. Esa cifra sube +1 por día laborable y cae de golpe cuando la orden sale.")
    linea("Por eso el Q1/Q3/promedio/máximo de la hoja 'Resumen' salen inflados: se calculan sobre los valores "
          "diarios de la orden más antigua, y basta una orden atascada para disparar el Q3 durante semanas.")
    linea("Todos los días son laborables del calendario de fábrica (no cuentan fines de semana ni festivos).")
    linea()

    linea("2. Cómo se calcula el LT de cada puesto", "seccion")
    linea("Cada caída de la antigüedad diaria es una 'salida': acaba de salir la orden más antigua y el valor del día "
          "anterior es su tiempo real en el puesto (fin − inicio, resolución ±1 día). Son lead times reales de "
          "órdenes terminadas, con sesgo conservador: representan a las órdenes NORMAL que más esperan.")
    linea(f"LT calculado = mediana (P50) de las salidas de los últimos {VENTANA_MESES} meses, redondeada hacia "
          "arriba (mínimo 1 día). Protección = P75 − LT. No se usa el máximo ni el Q3 por puesto (lanzan demasiado "
          "pronto, suben el WIP y alargan la cola: 'síndrome del lead time'), ni el promedio (lo distorsiona una "
          "sola orden atascada).")
    linea()

    linea("3. Cómo usarlo para lanzar las órdenes", "seccion")
    linea("La protección se pone una sola vez al final de la ruta, no en cada puesto:")
    linea("    Colchón de ruta = raíz cuadrada de la suma de (Protección de cada puesto)²   (redondeado hacia arriba)")
    linea("    Fecha de lanzamiento = Fecha de necesidad − Colchón − Suma de LT de los puestos de la ruta   (días laborables)")
    linea("    Fecha objetivo de cada operación = Fecha de necesidad − Colchón − Suma de LT de las operaciones posteriores")
    linea("La nueva pantalla de prioridades ordena por la fecha objetivo de la operación actual (o por la holgura = "
          "fecha objetivo − hoy). AOG sigue siempre por delante. Las órdenes de subconjuntos se encadenan: su fecha "
          "de necesidad es la fecha objetivo de la operación del padre que las consume.")
    linea()
    linea("Ejemplo con una ruta tipo (valores de este cálculo):", "negrita")
    linea("Puesto", "negrita", "LT", "Protección")
    suma_lt, suma_cuadrados, suma_p75 = 0, 0, 0
    for puesto in RUTA_EJEMPLO:
        r = por_puesto.get(puesto)
        if not r or r["lt"] is None:
            continue
        linea(f"    {puesto}", None, r["lt"], r["proteccion"])
        suma_lt += r["lt"]
        suma_cuadrados += r["proteccion"] ** 2
        suma_p75 += r["lt"] + r["proteccion"]
    colchon = math.ceil(math.sqrt(suma_cuadrados))
    linea(f"Suma de LT = {suma_lt} días; colchón = {colchon} días; lanzar {suma_lt + colchon} días laborables antes "
          f"de la necesidad (sumando el P75 de cada puesto saldrían {suma_p75}).", "negrita")
    linea()

    linea("4. Congelar y revisar", "seccion")
    linea("• El LT se congela en el ERP como parámetro por puesto con fecha de vigencia. Las órdenes ya lanzadas no se "
          "vuelven a fechar: el valor nuevo solo aplica a los lanzamientos nuevos.")
    linea(f"• Revisión mensual (primer día laborable) ejecutando este script con la exportación nueva y "
          f"--vigente con el fichero del mes anterior. Ventana móvil de {VENTANA_MESES} meses con esta exportación; "
          f"de {VENTANA_RECIENTE_MESES} meses cuando haya datos por orden.")
    linea(f"• Banda muerta: solo se cambia si la diferencia con el vigente es ≥ {BANDA_PCT:.0%} y ≥ {BANDA_MIN_DIAS} "
          f"día laborable. Cambio máximo por revisión: ±{TOPE_CAMBIO_PCT:.0%} (mínimo 1 día).")
    linea("• La columna 'LT propuesto (a congelar)' es la que se carga en el ERP. Si se corrige a mano (p. ej. para "
          "los puestos NO FIABLE), el mes siguiente el script la toma como vigente.")
    linea("• Revisión fuera de ciclo ante cambios estructurales: máquina nueva o dada de baja, cambio de turnos o de "
          "rutas, programa nuevo.")
    linea("• Tras arrancar la nueva pantalla de prioridades: revisión mensual al menos 6 meses (cambia la disciplina "
          "de cola y con ella el LT real). Después, trimestral en puestos estables y mensual en cuellos de botella.")
    linea()

    linea("5. KPI de control", "seccion")
    linea("% de operaciones que terminan dentro de su LT planificado, por puesto. Objetivo: 80–90 %. "
          "Por encima del 95 % sostenido: LT holgado, bajarlo. Por debajo del 70 %: antes de subirlo, revisar "
          "capacidad y bloqueos.")
    linea()

    linea("6. Avisos", "seccion")
    linea(f"• NO FIABLE: puesto de la lista manual del script, o menos de {NO_FIABLE_MAX_SALIDAS} salidas con la "
          f"orden más antigua por encima de {NO_FIABLE_ANTIGUEDAD} días. Sin LT calculado.")
    linea(f"• POCOS DATOS: menos de {MIN_SALIDAS} salidas en la ventana. Valor orientativo.")
    linea(f"• BLOQUEO ACTUAL: la orden más antigua lleva más de max({BLOQUEO_MIN_DIAS}, {BLOQUEO_FACTOR} × LT) días. "
          "El LT es válido; hay que desbloquear la orden.")
    linea(f"• CAMBIO RECIENTE: la mediana de los últimos {VENTANA_RECIENTE_MESES} meses difiere de la de "
          f"{VENTANA_MESES} meses en más de max({CAMBIO_MIN_DIAS} días, {CAMBIO_PCT:.0%}).")
    linea(f"• WIP CRECIENTE: el WIP de los últimos {WIP_DIAS_RECIENTES} días supera {WIP_FACTOR} × el WIP medio de "
          "la ventana. Si no se actúa sobre la capacidad, el LT real subirá.")
    linea()

    linea("7. Datos necesarios para la versión definitiva", "seccion")
    linea("Exportación del ERP con una fila por orden y operación: nº de lanzamiento, nº de operación, puesto, PN, "
          "cantidad, prioridad, fecha de llegada al puesto (fin de la operación anterior), fecha de fin de la "
          "operación e indicador de si estuvo retenida. Con ella se calculan P50 y P75 reales de todas las órdenes "
          "(no solo de la más antigua) y se puede medir directamente el KPI de cumplimiento.")


def hoja_evolucion(libro, resultados, areas, corte):
    hoja = libro.create_sheet("Evolución mensual")
    inicio = min(r["serie"][0][0] for r in resultados)
    meses = []
    mes = dt.date(inicio.year, inicio.month, 1)
    while mes <= corte:
        meses.append((mes.year, mes.month))
        mes = dt.date(mes.year + mes.month // 12, mes.month % 12 + 1, 1)
    etiquetas = [f"{MESES[m - 1]}-{a % 100:02d}" for a, m in meses]
    titulos = (["Área", "Puesto"] + [f"LT P50 {e}" for e in etiquetas]
               + [f"Nº salidas {e}" for e in etiquetas] + [f"WIP medio {e}" for e in etiquetas])
    cabecera(hoja, titulos, [14, 36] + [9] * (3 * len(meses)))

    for r in resultados:
        lt_mes = [percentil([v for f, v in r["salidas"] if (f.year, f.month) == am], 0.5) for am in meses]
        n_mes = [sum(1 for f, _ in r["salidas"] if (f.year, f.month) == am) for am in meses]
        wip_mes = [media([w for f, w, _ in r["serie"] if (f.year, f.month) == am]) for am in meses]
        hoja.append([areas[r["puesto"]], r["puesto"]] + lt_mes + n_mes + wip_mes)
        fila = hoja.max_row
        for i in range(len(meses)):
            hoja.cell(row=fila, column=3 + i).number_format = "0.0"
            hoja.cell(row=fila, column=3 + 2 * len(meses) + i).number_format = "0"
    hoja.auto_filter.ref = hoja.dimensions


# ---------------------------------------------------------------------------
def main():
    carpeta = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="Calcula el LT planificado por puesto de trabajo.")
    parser.add_argument("entrada", nargs="?", default=str(carpeta / "Resumen_LT_Año.xlsx"),
                        help="Excel exportado del ERP con la hoja 'Detalle' (por defecto Resumen_LT_Año.xlsx)")
    parser.add_argument("--vigente", help="LT_planificado_AAAA-MM.xlsx del mes anterior (LT congelado)")
    args = parser.parse_args()

    entrada = Path(args.entrada)
    libro_entrada = load_workbook(entrada, read_only=True, data_only=True)
    series, areas = leer_detalle(libro_entrada)
    q3_resumen = leer_resumen_q3(libro_entrada)
    libro_entrada.close()
    if not series:
        sys.exit("La hoja 'Detalle' no tiene datos.")

    corte = max(serie[-1][0] for serie in series.values())
    ini6 = primer_dia_mes(corte, VENTANA_MESES - 1)
    ini3 = primer_dia_mes(corte, VENTANA_RECIENTE_MESES - 1)
    salida = entrada.parent / f"LT_planificado_{corte:%Y-%m}.xlsx"
    if args.vigente and Path(args.vigente).resolve() == salida.resolve():
        sys.exit(f"El fichero vigente y el de salida son el mismo ({salida.name}): "
                 "usa como --vigente el fichero del mes anterior.")
    vigentes = leer_vigente(args.vigente) if args.vigente else {}

    resultados = [calcular_puesto(p, series[p], corte, ini6, ini3, q3_resumen)
                  for p in sorted(series, key=lambda p: (areas[p], p))]

    libro = Workbook()
    hoja_listado(libro, resultados, areas, vigentes, bool(args.vigente))
    hoja_criterio(libro, resultados, corte, ini6, ini3, Path(args.vigente).name if args.vigente else None)
    hoja_evolucion(libro, resultados, areas, corte)
    libro.save(salida)

    print(f"Fecha de corte: {corte:%d/%m/%Y} | ventana: {ini6:%d/%m/%Y} - {corte:%d/%m/%Y}")
    print(f"Puestos: {len(resultados)} | con LT calculado: {sum(r['lt'] is not None for r in resultados)}")
    for tipo in ("NO FIABLE", "POCOS DATOS", "BLOQUEO ACTUAL", "CAMBIO RECIENTE", "WIP CRECIENTE"):
        puestos = [r["puesto"] for r in resultados if tipo in r["aviso"]]
        print(f"  {tipo} ({len(puestos)}): {', '.join(puestos) or '-'}")
    print(f"Generado: {salida}")


if __name__ == "__main__":
    main()
